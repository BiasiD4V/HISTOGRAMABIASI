from __future__ import annotations
import json, re
from pathlib import Path
from typing import Any, Dict, List, Optional

DATA_PATH = Path(__file__).parent / "data" / "cost_base_biasi.json"

ROLE_ALIASES = {
    "Eletricista": "ELETRICISTA MONTADOR",
    "Eletricista Montador": "ELETRICISTA MONTADOR",
    "Eletricista Força e Controle": "ELETRICISTA FORÇA E CONTROLE",
    "Ajudante de Eletricista": "AJUDANTE DE ELETRICISTA",
    "Encanador": "ENCANADOR B",
    "Encanador B": "ENCANADOR B",
    "Ajudante de Hidráulica": "AJUDANTE DE HIDRAULICA",
    "Ajudante de Hidraulica": "AJUDANTE DE HIDRAULICA",
    "Ajudante Geral": "AJUDANTE GERAL",
    "Soldador/Serralheiro": "SERRALHEIRO",
    "Serralheiro": "SERRALHEIRO",
    "Pedreiro": "PEDREIRO",
    "Instrumentista": "ELETRICISTA FORÇA E CONTROLE",
    "Téc. Automação": "ELETRICISTA FORÇA E CONTROLE",
    "Técnico de Automação": "ELETRICISTA FORÇA E CONTROLE",
    "Encarregado Elétrico": "ENCARREGADO DE ELETRICA PREDIAL",
    "Encarregado de Elétrica": "ENCARREGADO DE ELETRICA PREDIAL",
    "Encarregado Hidráulico": "ENCARREGADO DE HIDRAULICA",
    "Supervisor": "SUPERVIROR DE OBRAS(DI)",
    "Supervisor de Obras": "SUPERVIROR DE OBRAS(DI)",
    "Engenheiro": "ENGENHEIRO ELETRICISTA 6H",
    "Engenheiro Eletricista": "ENGENHEIRO ELETRICISTA 6H",
    "Téc. Segurança": "TEC. SEG. TRABALHO (DI)",
    "Técnico Segurança": "TEC. SEG. TRABALHO (DI)",
    "Almoxarife": "ALMOXARIFE",
}

def _norm(s: str) -> str:
    s = (s or "").strip().upper()
    s = re.sub(r"\s+", " ", s)
    return s

def load_default_costs() -> Dict[str, float]:
    return json.loads(DATA_PATH.read_text(encoding="utf-8"))

def import_costs_from_xlsx(path: str | Path) -> Dict[str, float]:
    """Importa a aba CUSTOS E PROVISÕES. Usa coluna FUNÇÃO e TOTAL DE DESPESAS MENSAL."""
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb["CUSTOS E PROVISÕES"] if "CUSTOS E PROVISÕES" in wb.sheetnames else wb[wb.sheetnames[0]]
    header_row = None
    function_col = total_col = None
    for r in range(1, min(ws.max_row, 20) + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        for c, v in enumerate(vals, start=1):
            vv = _norm(str(v)) if v is not None else ""
            if vv == "FUNÇÃO" or vv == "FUNCAO":
                header_row = r; function_col = c
            if "TOTAL DE DESPESAS MENSAL" in vv:
                header_row = r; total_col = c
        if function_col and total_col:
            break
    if not function_col or not total_col:
        raise ValueError("Não encontrei as colunas FUNÇÃO e TOTAL DE DESPESAS MENSAL na planilha de custos.")
    out: Dict[str, float] = {}
    for r in range(header_row + 1, ws.max_row + 1):
        name = ws.cell(r, function_col).value
        value = ws.cell(r, total_col).value
        if isinstance(name, str) and isinstance(value, (int, float)) and value > 0:
            out[_norm(name)] = round(float(value), 2)
    if not out:
        raise ValueError("A planilha de custos não trouxe funções válidas.")
    return out

def normalize_role(role: str) -> str:
    if role in ROLE_ALIASES:
        return ROLE_ALIASES[role]
    n = _norm(role)
    # match loose aliases
    if "AJUD" in n and "ELE" in n: return "AJUDANTE DE ELETRICISTA"
    if "AJUD" in n and ("HID" in n or "TUB" in n): return "AJUDANTE DE HIDRAULICA"
    if "ELETRICISTA" in n and "FOR" in n: return "ELETRICISTA FORÇA E CONTROLE"
    if "ELETRICISTA" in n: return "ELETRICISTA MONTADOR"
    if "ENCAN" in n: return "ENCANADOR B"
    if "SERR" in n or "SOLD" in n: return "SERRALHEIRO"
    if "PEDR" in n: return "PEDREIRO"
    if "ENCAR" in n and "HID" in n: return "ENCARREGADO DE HIDRAULICA"
    if "ENCAR" in n and "ELE" in n: return "ENCARREGADO DE ELETRICA PREDIAL"
    if "SEG" in n: return "TEC. SEG. TRABALHO (DI)"
    if "SUP" in n: return "SUPERVIROR DE OBRAS(DI)"
    if "ENG" in n and "CIV" in n: return "ENGENHEIRO CIVIL 6H"
    if "ENG" in n: return "ENGENHEIRO ELETRICISTA 6H"
    if "ALMOX" in n: return "ALMOXARIFE"
    if "AJUD" in n: return "AJUDANTE GERAL"
    return n

def calculate_costs(
    weekly_histogram: List[Dict[str, Any]],
    dias_semana: int,
    jornada_h_dia: float,
    costs: Optional[Dict[str, float]] = None,
) -> Dict[str, Any]:
    costs = costs or load_default_costs()
    weekly = []
    total = 0.0
    by_role: Dict[str, float] = {}
    missing_roles = set()
    for row in weekly_histogram:
        row_cost = 0.0
        detail = {}
        for role, qtd in row.get("funcoes", {}).items():
            qtd = float(qtd or 0)
            if qtd <= 0:
                continue
            role_key = normalize_role(role)
            monthly = float(costs.get(role_key, 0.0))
            if not monthly:
                missing_roles.add(role_key)
            weekly_cost = monthly / 22.0 * float(dias_semana) * (float(jornada_h_dia) / 8.0) * qtd
            detail[role] = {
                "qtd": qtd,
                "funcao_base": role_key,
                "custo_mensal": round(monthly, 2),
                "custo_semanal": round(weekly_cost, 2),
                "formula": "custo_mensal/22*dias_semana*(jornada_h_dia/8)*qtd",
            }
            row_cost += weekly_cost
            by_role[role] = by_role.get(role, 0.0) + weekly_cost
        weekly.append({"semana": row.get("semana"), "custo_mo": round(row_cost, 2), "detalhe": detail})
        total += row_cost
    return {
        "total_mo": round(total, 2),
        "weekly_costs": weekly,
        "by_role": {k: round(v, 2) for k, v in by_role.items()},
        "missing_roles": sorted(missing_roles),
        "cost_formula": "custo semanal = custo mensal/22*dias_semana*(jornada_h_dia/8)*quantidade",
        "cost_base": costs,
    }
