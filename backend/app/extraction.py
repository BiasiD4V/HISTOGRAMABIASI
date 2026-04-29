from __future__ import annotations
import asyncio, io
from pathlib import Path
from typing import AsyncGenerator, List, Tuple
from fastapi import UploadFile
from .schemas import DocumentText

MAX_CHARS_PER_FILE = 70_000
MAX_PDF_PAGES = 10
EXTRACTION_TIMEOUT_SECONDS = 35

def _trim(text: str) -> str:
    text = text or ""
    if len(text) > MAX_CHARS_PER_FILE:
        return text[:MAX_CHARS_PER_FILE] + "\n\n[TRUNCADO PARA PROCESSAMENTO]"
    return text

async def extract_uploads(files: List[UploadFile]) -> List[DocumentText]:
    docs: List[DocumentText] = []
    async for _event, doc in extract_uploads_progress(files):
        if doc is not None:
            docs.append(doc)
    return docs


async def extract_raw_uploads_progress(items: List[tuple[str, bytes]]) -> AsyncGenerator[Tuple[dict, DocumentText | None], None]:
    items = items or []
    total = len(items)
    yield {"type":"status", "message":f"Recebi {total} arquivo(s). Iniciando extração..."}, None
    for idx, (name, raw) in enumerate(items, start=1):
        name = name or f"arquivo_{idx}"
        yield {"type":"status", "message":f"Extraindo {idx}/{total}: {name}"}, None
        try:
            text, status = await asyncio.wait_for(
                asyncio.to_thread(_extract_raw, name, raw),
                timeout=EXTRACTION_TIMEOUT_SECONDS,
            )
            doc = DocumentText(filename=name, status=status, text=_trim(text), chars=len(text))
            yield {"type":"status", "message":f"OK {idx}/{total}: {name} ({doc.chars} caracteres)"}, doc
        except asyncio.TimeoutError:
            doc = DocumentText(filename=name, status="parcial", text=f"[Extração interrompida por tempo limite: {name}. O arquivo foi usado apenas como referência nominal.]", chars=0, error="timeout")
            yield {"type":"status", "message":f"Tempo limite em {name}; seguindo com os demais."}, doc
        except Exception as exc:
            doc = DocumentText(filename=name, status="erro", text=f"[Falha ao extrair {name}]", chars=0, error=str(exc))
            yield {"type":"status", "message":f"Falha em {name}: {exc}"}, doc
    yield {"type":"status", "message":"Extração concluída. Iniciando agentes..."}, None

async def extract_uploads_progress(files: List[UploadFile]) -> AsyncGenerator[Tuple[dict, DocumentText | None], None]:
    files = files or []
    total = len(files)
    yield {"type":"status", "message":f"Recebi {total} arquivo(s). Iniciando extração..."}, None
    for idx, file in enumerate(files, start=1):
        name = file.filename or f"arquivo_{idx}"
        yield {"type":"status", "message":f"Extraindo {idx}/{total}: {name}"}, None
        try:
            raw = await file.read()
            text, status = await asyncio.wait_for(
                asyncio.to_thread(_extract_raw, name, raw),
                timeout=EXTRACTION_TIMEOUT_SECONDS,
            )
            doc = DocumentText(filename=name, status=status, text=_trim(text), chars=len(text))
            yield {"type":"status", "message":f"OK {idx}/{total}: {name} ({doc.chars} caracteres)"}, doc
        except asyncio.TimeoutError:
            doc = DocumentText(filename=name, status="parcial", text=f"[Extração interrompida por tempo limite: {name}. O arquivo foi usado apenas como referência nominal.]", chars=0, error="timeout")
            yield {"type":"status", "message":f"Tempo limite em {name}; seguindo com os demais."}, doc
        except Exception as exc:
            doc = DocumentText(filename=name, status="erro", text=f"[Falha ao extrair {name}]", chars=0, error=str(exc))
            yield {"type":"status", "message":f"Falha em {name}: {exc}"}, doc
    yield {"type":"status", "message":"Extração concluída. Iniciando agentes..."}, None

def _extract_raw(name: str, raw: bytes) -> tuple[str, str]:
    lower = name.lower()
    if lower.endswith((".txt", ".csv", ".md", ".json")):
        return raw.decode("utf-8", errors="ignore"), "ok"
    if lower.endswith(".pdf"):
        return _extract_pdf(raw), "ok"
    if lower.endswith(".docx"):
        return _extract_docx(raw), "ok"
    if lower.endswith((".xlsx", ".xlsm", ".xls")):
        return _extract_xlsx(raw), "ok"
    return f"[Arquivo anexado, mas tipo não extraído automaticamente: {name}]", "parcial"

def _extract_pdf(raw: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(raw), strict=False)
    pages = []
    for p in reader.pages[:MAX_PDF_PAGES]:
        try:
            pages.append(p.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n".join(pages)

def _extract_docx(raw: bytes) -> str:
    from docx import Document
    doc = Document(io.BytesIO(raw))
    blocks = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            vals = [c.text.strip() for c in row.cells]
            if any(vals):
                blocks.append(" | ".join(vals))
    return "\n".join(blocks)

def _extract_xlsx(raw: bytes) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(raw), data_only=True, read_only=True)
    chunks = []
    for ws in wb.worksheets[:8]:
        chunks.append(f"### PLANILHA: {ws.title}")
        max_rows = min(ws.max_row, 160)
        max_cols = min(ws.max_column, 45)
        for row in ws.iter_rows(min_row=1, max_row=max_rows, max_col=max_cols, values_only=True):
            vals = [str(v).strip() if v is not None else "" for v in row]
            if any(vals):
                chunks.append(";".join(vals))
    return "\n".join(chunks)
